# -*- encoding: utf-8 -*-
"""Journal d'activité : export puis vidage (voir apps/views.py,
_exporter_et_vider_journal, et apps/gtc_data.py, vider_journal).

Régression : le vidage semblait ne jamais s'exécuter alors que le
fichier était bien téléchargé. La cause réelle n'était PAS côté serveur
(déjà correct : export d'abord, vidage seulement si l'export réussit) —
c'était l'absence de rafraîchissement de la page après une réponse de
succès qui EST le fichier exporté (voir apps/templates/pages/journal.html).
Ces tests couvrent la partie serveur : contenu réel de l'export, état
de la table juste après, et les deux protections (rien vidé si l'export
échoue, message clair si le vidage échoue après un export réussi)."""
import io

from openpyxl import load_workbook

from apps import db
from apps.models import JournalActivite


def _peupler_journal(admin_id, n=5):
    for i in range(n):
        db.session.add(JournalActivite(
            horodatage=__import__("datetime").datetime.now(),
            utilisateur_id=admin_id, identifiant="admin.test",
            action="mouvement_test", description=f"Entrée de test numéro {i}",
        ))
    db.session.commit()


def test_export_excel_puis_vidage_fichier_complet_table_vide_apres(client, login, admin):
    _peupler_journal(admin, n=5)
    login("admin.test")
    with client.session_transaction() as s:
        pass  # la connexion elle-même journalise une entrée "connexion"

    avant = JournalActivite.query.count()
    assert avant == 6  # 5 + la connexion

    resp = client.post("/pages/journal/exporter-et-vider.xlsx")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers.get("Content-Type", "")

    classeur = load_workbook(io.BytesIO(resp.data))
    descriptions = [row[3].value for feuille in [classeur["Journal"]]
                    for row in feuille.iter_rows(min_row=2, values_only=False)]
    for i in range(5):
        assert f"Entrée de test numéro {i}" in descriptions
    assert len(descriptions) == avant  # le fichier contient TOUTES les entrées

    lignes_restantes = JournalActivite.query.all()
    assert len(lignes_restantes) == 1  # table vide à part la trace de purge
    assert lignes_restantes[0].action == "purge_journal"


def test_export_pdf_puis_vidage(client, login, admin):
    _peupler_journal(admin, n=3)
    login("admin.test")
    resp = client.post("/pages/journal/exporter-et-vider.pdf")
    assert resp.status_code == 200
    assert resp.headers.get("Content-Type") == "application/pdf"
    assert resp.data.startswith(b"%PDF")

    lignes_restantes = JournalActivite.query.all()
    assert len(lignes_restantes) == 1
    assert lignes_restantes[0].action == "purge_journal"


def test_journal_deja_vide_ni_export_ni_vidage(client, login, admin):
    login("admin.test")
    JournalActivite.query.delete()
    db.session.commit()

    resp = client.post("/pages/journal/exporter-et-vider.pdf", follow_redirects=True)
    assert "déjà vide" in resp.get_data(as_text=True)
    assert JournalActivite.query.count() == 0


def test_echec_export_ne_vide_pas_le_journal(client, login, admin, monkeypatch):
    """Protection n°1 : si la génération de l'export échoue, rien n'est
    supprimé (jamais de perte de données sans export réussi)."""
    _peupler_journal(admin, n=5)
    login("admin.test")
    avant = JournalActivite.query.count()

    import apps.views as views_module

    def _panne(*args, **kwargs):
        raise RuntimeError("panne simulée de la génération d'export")

    monkeypatch.setattr(views_module.exports, "journal_excel", _panne)
    resp = client.post("/pages/journal/exporter-et-vider.xlsx", follow_redirects=True)

    page = resp.get_data(as_text=True)
    assert "n&#39;a PAS" in page or "n'a PAS" in page
    assert JournalActivite.query.count() == avant  # intact


def test_echec_vidage_apres_export_reussi_message_clair(client, login, admin, monkeypatch):
    """Protection n°2 : si le vidage échoue APRÈS un export déjà réussi,
    un message d'erreur clair est affiché — jamais un faux succès."""
    _peupler_journal(admin, n=5)
    login("admin.test")
    avant = JournalActivite.query.count()

    import apps.views as views_module

    def _panne(acteur):
        raise RuntimeError("panne simulée du vidage")

    monkeypatch.setattr(views_module, "vider_journal", _panne)
    resp = client.post("/pages/journal/exporter-et-vider.pdf", follow_redirects=True)

    page = resp.get_data(as_text=True)
    assert "aucune entrée n&#39;a été supprimée" in page or "aucune entrée n'a été supprimée" in page
    assert JournalActivite.query.count() == avant  # intact, aucune ligne perdue


def test_formulaire_journal_intercepte_en_js_pour_rafraichir_la_page(client, login, admin):
    """Non-régression du vrai problème signalé : la page doit maintenant
    forcer un rechargement après l'action (voir gtc-stock côté client),
    au lieu de laisser un téléchargement de fichier sans navigation
    faire croire que rien ne s'est passé."""
    _peupler_journal(admin, n=1)
    login("admin.test")
    resp = client.get("/pages/journal/")
    page = resp.get_data(as_text=True)
    assert 'id="form-exporter-vider-journal"' in page
    assert "window.location.reload" in page
