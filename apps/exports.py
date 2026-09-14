# -*- encoding: utf-8 -*-
"""
GTC Stock — Export PDF et Excel de la fiche de stock et du rapprochement.

Deux formats, mêmes données que ce que la page affiche à l'écran : PDF
pour imprimer/partager en lecture seule, Excel pour retravailler les
chiffres (trier, filtrer, recalculer) dans un tableur.
"""

import io
from datetime import datetime
from xml.sax.saxutils import escape as _echapper_xml

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

_STYLES = getSampleStyleSheet()
_STYLE_CELLULE_JOURNAL = ParagraphStyle(
    "CelluleJournal", parent=_STYLES["Normal"], fontSize=8, leading=10,
)


def _horodatage():
    return datetime.now().strftime("%d/%m/%Y à %H:%M")


# ---------------------------------------------------------------------
# PDF — mise en page commune (titre + sous-titre + un ou plusieurs tableaux)
# ---------------------------------------------------------------------
def _construire_pdf(titre, sous_titre, tableaux):
    """tableaux : liste de (legende_ou_None, liste_de_lignes[, largeurs_colonnes])
    — le 3e élément est optionnel (None par défaut = largeurs automatiques,
    utilisé quand une colonne doit être forcée, ex. une description longue
    mise en forme avec Paragraph pour l'export du journal). La première
    ligne de chaque liste est traitée comme l'en-tête."""
    tampon = io.BytesIO()
    doc = SimpleDocTemplate(
        tampon, pagesize=A4,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    elements = [
        Paragraph(titre, _STYLES["Title"]),
        Paragraph(sous_titre, _STYLES["Normal"]),
        Spacer(1, 0.5 * cm),
    ]
    for item in tableaux:
        legende, lignes = item[0], item[1]
        largeurs_colonnes = item[2] if len(item) > 2 else None
        if legende:
            elements.append(Paragraph(legende, _STYLES["Heading3"]))
        table = Table(lignes, repeatRows=1, colWidths=largeurs_colonnes)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1b2559")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d0d5dd")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f8fa")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 0.5 * cm))
    doc.build(elements)
    return tampon.getvalue()


# ---------------------------------------------------------------------
# Excel — feuille unique, en-tête en gras
# ---------------------------------------------------------------------
def _construire_excel(nom_feuille, entetes, lignes):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = nom_feuille
    feuille.append(entetes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for ligne in lignes:
        feuille.append(ligne)
    for i, entete in enumerate(entetes, start=1):
        feuille.column_dimensions[get_column_letter(i)].width = max(14, len(str(entete)) + 4)

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------------
# Fiche de stock
# ---------------------------------------------------------------------
def fiche_stock_pdf(article, historique):
    infos = [
        ["Référence", article.reference],
        ["Quantité actuelle", str(article.quantite)],
        ["Seuil d'alerte", str(article.seuil)],
        ["Statut", article.statut],
        ["Dernier mouvement", article.dernier_mouvement or "—"],
    ]
    # h["detail"] retombe sur le motif quand une Entree n'a pas de
    # référence externe (ex. régularisation, sans bordereau) — voir
    # gtc_data._mouvement_vers_dict.
    lignes_historique = [["Date", "Type", "Référence", "Quantité", "Solde après mouvement"]]
    lignes_historique += [
        [h["date"], h["type"], h["detail"], h["quantite"], str(h["solde"])]
        for h in historique
    ] or [["Aucun mouvement enregistré.", "", "", "", ""]]

    return _construire_pdf(
        f"Fiche de stock — {article.nom}",
        f"Généré le {_horodatage()} — GTC Stock",
        [(None, infos), ("Historique des mouvements", lignes_historique)],
    )


def fiche_stock_excel(article, historique):
    classeur = Workbook()

    resume = classeur.active
    resume.title = "Article"
    resume.append(["Champ", "Valeur"])
    for cellule in resume[1]:
        cellule.font = Font(bold=True)
    for champ, valeur in [
        ("Nom", article.nom), ("Référence", article.reference),
        ("Quantité actuelle", article.quantite), ("Seuil d'alerte", article.seuil),
        ("Statut", article.statut), ("Dernier mouvement", article.dernier_mouvement or "—"),
    ]:
        resume.append([champ, valeur])
    resume.column_dimensions["A"].width = 20
    resume.column_dimensions["B"].width = 30

    historique_feuille = classeur.create_sheet("Historique")
    historique_feuille.append(["Date", "Type", "Référence", "Quantité", "Solde après mouvement"])
    for cellule in historique_feuille[1]:
        cellule.font = Font(bold=True)
    for h in historique:
        historique_feuille.append([h["date"], h["type"], h["detail"], h["quantite_brute"], h["solde"]])
    for i in range(1, 6):
        historique_feuille.column_dimensions[get_column_letter(i)].width = 18

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------------
# Rapprochement Sage 100
# ---------------------------------------------------------------------
def rapprochement_pdf(lignes, sage_connecte):
    source = "Sage 100 (connexion en direct)" if sage_connecte else "données de démonstration (Sage 100 non connecté)"
    tableau = [["Article", "Quantité application", "Quantité Sage 100", "Écart", "Statut"]]
    tableau += [
        [l["article"], str(l["qte_app"]), str(l["qte_sage"]), str(l["ecart"]),
         "Conforme" if l["conforme"] else "Écart détecté"]
        for l in lignes
    ] or [["Aucune donnée de rapprochement.", "", "", "", ""]]

    return _construire_pdf(
        "Rapprochement comptable — GTC Stock",
        f"Généré le {_horodatage()} — Source : {source}",
        [(None, tableau)],
    )


def rapprochement_excel(lignes, sage_connecte):
    entetes = ["Article", "Quantité application", "Quantité Sage 100", "Écart", "Conforme"]
    donnees = [
        [l["article"], l["qte_app"], l["qte_sage"], l["ecart"], "Oui" if l["conforme"] else "Non"]
        for l in lignes
    ]
    return _construire_excel("Rapprochement", entetes, donnees)


# ---------------------------------------------------------------------
# Journal d'activité (audit log) — export complet avant purge (voir
# gtc_data.vider_journal) ou simple consultation imprimable.
# ---------------------------------------------------------------------
def journal_pdf(journal):
    """La colonne Description est mise en forme avec Paragraph (et des
    largeurs de colonnes fixes) pour ne pas être tronquée : contrairement
    aux autres exports, ces descriptions peuvent être assez longues."""
    lignes = [["Horodatage", "Identifiant", "Action", "Description"]]
    lignes += [
        [
            j.horodatage.strftime("%d/%m/%Y %H:%M"),
            j.identifiant or "—",
            j.action_libelle,
            Paragraph(_echapper_xml(j.description), _STYLE_CELLULE_JOURNAL),
        ]
        for j in journal
    ] or [["Aucune entrée.", "", "", ""]]

    return _construire_pdf(
        "Journal d'activité — GTC Stock",
        f"Export complet généré le {_horodatage()} — {len(journal)} entrée(s)",
        [(None, lignes, [3.3 * cm, 3 * cm, 3.3 * cm, 8.2 * cm])],
    )


def journal_excel(journal):
    entetes = ["Horodatage", "Identifiant", "Action", "Description"]
    donnees = [
        [j.horodatage.strftime("%d/%m/%Y %H:%M"), j.identifiant or "—", j.action_libelle, j.description]
        for j in journal
    ]
    return _construire_excel("Journal", entetes, donnees)
